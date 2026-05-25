"""Test parse case: **Câu N.** + đáp án in đậm trên dòng A/B/C/D."""

from __future__ import annotations

import unittest
from types import SimpleNamespace

from back.tools import (
    RE_OPT_LINE,
    RE_Q_PREFIX,
    any_bold_in_range,
    clean_question_text,
    detect_markdown_wrapped_answer_key,
    detect_mcq_answer_from_paragraph,
    match_question_prefix,
    normalize_line_for_parse,
    parse_option_line,
    paragraph_spans,
    question_number_from_line,
    RunSpan,
)


def _para(text: str, bold_ranges: list[tuple[int, int]] | None = None):
    """Mock paragraph: bold_ranges = [(start, end), ...] trên plain text."""
    bold_ranges = bold_ranges or []
    runs = []
    if not bold_ranges:
        runs.append(SimpleNamespace(text=text, bold=False))
    else:
        pos = 0
        bold_set = set()
        for s, e in bold_ranges:
            for i in range(s, e):
                bold_set.add(i)
        i = 0
        while i < len(text):
            is_bold = i in bold_set
            j = i + 1
            while j < len(text) and (j in bold_set) == is_bold:
                j += 1
            runs.append(SimpleNamespace(text=text[i:j], bold=is_bold))
            i = j
    return SimpleNamespace(text=text, runs=runs)


class TestMarkdownQuestionPrefix(unittest.TestCase):
    def test_cau_10_markdown(self):
        line = "**Câu 10.** Đâu là tag tạo ra chữ in đậm?"
        n = normalize_line_for_parse(line)
        self.assertTrue(RE_Q_PREFIX.match(n))
        self.assertEqual(question_number_from_line(line), 10)
        m = match_question_prefix(line)
        self.assertIsNotNone(m)
        n2 = normalize_line_for_parse(line)
        m2 = RE_Q_PREFIX.match(n2)
        q = clean_question_text(n2[m2.end() :])
        self.assertIn("tag tạo ra chữ in đậm", q)

    def test_cau_11_markdown(self):
        line = "**Câu 11.** Đâu là tag tạo ra màu nền của web?"
        self.assertEqual(question_number_from_line(line), 11)


class TestOptionLines(unittest.TestCase):
    def test_option_b_tag(self):
        line = "A. `<b>`"
        m = RE_OPT_LINE.match(normalize_line_for_parse(line))
        self.assertIsNotNone(m)
        self.assertEqual(m.group(1).upper(), "A")
        self.assertIn("b", m.group(3).lower())

    def test_option_body_bgcolor(self):
        line = 'D. `<body background-color="yellow">`'
        m = RE_OPT_LINE.match(line)
        self.assertIsNotNone(m)
        self.assertEqual(m.group(1).upper(), "D")

    def test_option_with_note(self):
        line = (
            'D. `<body background-color="yellow">` '
            "*(Lưu ý: Thuộc tính chuẩn trong CSS inline là style=\"background-color:yellow\".)*"
        )
        m = RE_OPT_LINE.match(line)
        self.assertIsNotNone(m)
        self.assertIn("Lưu ý", m.group(3))


class TestMarkdownWrappedAnswer(unittest.TestCase):
    """Đáp án **D. ...** — literal asterisk, không in đậm Word."""

    def test_cau_29_style_answer(self):
        line = "**D. .htm và .html**"
        self.assertEqual(detect_markdown_wrapped_answer_key(line), "D")

    def test_parse_option_line_md(self):
        parsed = parse_option_line("**D. .htm và .html**")
        self.assertIsNotNone(parsed)
        key, text, ans = parsed
        self.assertEqual(key, "D")
        self.assertEqual(text, ".htm và .html")
        self.assertEqual(ans, "D")

    def test_detect_without_bold(self):
        p = _para("**D. .htm và .html**")
        self.assertEqual(detect_mcq_answer_from_paragraph(p), "D")

    def test_question_line_not_answer(self):
        self.assertEqual(detect_markdown_wrapped_answer_key("**Câu 29.** Phần mở rộng"), "")

    def test_plain_option_no_false_positive(self):
        self.assertEqual(detect_markdown_wrapped_answer_key("A. .htm"), "")


class TestBoldAnswerDetection(unittest.TestCase):
    def test_bold_on_option_content(self):
        # A. <b> — chỉ phần <b> in đậm
        text = "A. `<b>`"
        bold_start = text.index("<")
        p = _para(text, [(bold_start, len(text))])
        self.assertEqual(detect_mcq_answer_from_paragraph(p), "A")

    def test_bold_on_option_d(self):
        text = 'D. `<body background-color="yellow">`'
        bold_start = text.index("<")
        p = _para(text, [(bold_start, len(text))])
        self.assertEqual(detect_mcq_answer_from_paragraph(p), "D")

    def test_no_bold(self):
        p = _para("B. bold")
        self.assertEqual(detect_mcq_answer_from_paragraph(p), "")

    def test_question_line_not_option(self):
        p = _para("**Câu 10.** Đâu là tag tạo ra chữ in đậm?", [(0, 9)])
        self.assertEqual(detect_mcq_answer_from_paragraph(p), "")


class TestParagraphSpans(unittest.TestCase):
    def test_any_bold_in_range(self):
        spans = [RunSpan(0, 3, False), RunSpan(3, 8, True)]
        self.assertTrue(any_bold_in_range(spans, 4, 7))
        self.assertFalse(any_bold_in_range(spans, 0, 3))


class TestParseDefaultIntegration(unittest.TestCase):
    """Mô phỏng luồng parse_default với vài đoạn văn."""

    def test_two_questions_flow(self):
        from back.tools import parse_default
        from docx import Document

        doc = Document()
        doc.add_paragraph("**Câu 10.** Đâu là tag tạo ra chữ in đậm?")
        p_a = doc.add_paragraph("A. `<b>`")
        for r in p_a.runs:
            if "<b>" in (r.text or ""):
                r.bold = True
        doc.add_paragraph("B. bold")
        doc.add_paragraph("C. `<bld>`")
        doc.add_paragraph("D. `<bb>`")

        doc.add_paragraph("**Câu 11.** Đâu là tag tạo ra màu nền của web?")
        doc.add_paragraph('A. `<body color="yellow">`')
        doc.add_paragraph('B. `<body bgcolor="yellow">`')
        doc.add_paragraph("C. `<background>yellow</background>`")
        p_d = doc.add_paragraph('D. `<body background-color="yellow">`')
        for r in p_d.runs:
            if "background-color" in (r.text or ""):
                r.bold = True

        quiz = parse_default(doc)
        self.assertGreaterEqual(len(quiz), 2)
        q10 = next((x for x in quiz if "in đậm" in x["q"]), None)
        self.assertIsNotNone(q10)
        self.assertEqual(q10["ans"], "A")
        q11 = next((x for x in quiz if "màu nền" in x["q"]), None)
        self.assertIsNotNone(q11)
        self.assertEqual(q11["ans"], "D")

    def test_cau_29_markdown_answer(self):
        from back.tools import parse_default
        from docx import Document

        doc = Document()
        doc.add_paragraph("**Câu 29.** Phần mở rộng của tập tin HTML là?")
        doc.add_paragraph("A. .htm")
        doc.add_paragraph("B. .html")
        doc.add_paragraph("C. xhtml")
        doc.add_paragraph("**D. .htm và .html**")

        quiz = parse_default(doc)
        q29 = next((x for x in quiz if "mở rộng" in x["q"]), None)
        self.assertIsNotNone(q29)
        self.assertEqual(q29["ans"], "D")
        self.assertEqual(q29["opts"]["D"], ".htm và .html")


if __name__ == "__main__":
    unittest.main()
